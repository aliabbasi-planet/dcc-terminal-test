from __future__ import annotations

from io import BytesIO

import pandas as pd
import pytest
from openpyxl import Workbook

from dcc_console.remediation import DEFAULT_OBJECTS, batch_upload
from dcc_console.remediation import worklist as wl
from dcc_console.remediation.batch_upload import parse_workbook
from dcc_console.remediation.mapping import TRACKED_CHECK_COLUMNS
from dcc_console.remediation.tab import _load_batch_targets, _uploaded_batch_rows
from dcc_console.remediation.worklist import WorklistFilters

HEADERS = [
    "Change Type",
    "instance_identifier",
    "location_no",
    "location_desc",
    "country_name",
    "bank_merchant_id",
    "acquirer_name",
    "TerminalID",
]


def _workbook_bytes(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_parse_workbook_maps_handler_change_and_normalises_location_padding():
    content = _workbook_bytes(
        [[
            "handler_dccEnableNfc",
            "I0001",
            "011276",
            "Shop",
            "Spain",
            "00123",
            "Planet",
            "00012345",
        ]]
    )

    [target] = parse_workbook(content)

    assert target.terminal_identifier == "00012345"
    assert target.check_column == "HANDLER_DCCENABLENFC_CHECK_O"
    assert target.metadata["LOCATION_NO"] == "11276"
    assert target.metadata["LOCATION_DESC"] == "Shop"
    assert target.metadata["BANK_MERCHANT_ID"] == "00123"
    assert target.metadata["COUNTRY_NAME"] == "Spain"


def test_parse_workbook_accepts_country_header_alias():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Change Type", "instance_identifier", "location_no", "location_desc",
                  "country", "bank_merchant_id", "acquirer_name", "TerminalID"])
    sheet.append(["handler_dccEnableNfc", "I0001", "11276", "Shop", "Spain", "00123",
                  "Planet", "T-123"])
    output = BytesIO()
    workbook.save(output)

    [target] = parse_workbook(output.getvalue())

    assert target.metadata["COUNTRY_NAME"] == "Spain"


def test_parse_workbook_uses_first_column_as_check_and_maps_other_columns_by_name():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Any heading", "TerminalID", "location_desc", "instance_identifier"])
    sheet.append(["HANDLER_DCCENABLENFC_CHECK_O", "T-123", "Shop", "I-001"])
    output = BytesIO()
    workbook.save(output)

    [target] = parse_workbook(output.getvalue())

    assert target.check_column == "HANDLER_DCCENABLENFC_CHECK_O"
    assert target.terminal_identifier == "T-123"
    assert target.metadata == {"LOCATION_DESC": "Shop", "INSTANCE_IDENTIFIER": "I-001"}


def test_parse_workbook_maps_dcc_provider_code():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Check", "TerminalID", "dccprovider_code"])
    sheet.append(["dccEnable", "T-123", "PLANETPAYMENT"])
    output = BytesIO()
    workbook.save(output)

    [target] = parse_workbook(output.getvalue())

    assert target.metadata["DCCPROVIDER_CODE"] == "PLANETPAYMENT"


def test_parse_workbook_maps_cortex_printout_fields():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([
        "Check", "TerminalID", "printout_type_receipt", "printout_type_template_dcc",
        "printout_type_display",
    ])
    sheet.append(["Receipt template DCC", "T-123", "DCC", "PLANET_PAX", "Customer"])
    output = BytesIO()
    workbook.save(output)

    [target] = parse_workbook(output.getvalue())

    assert target.metadata == {
        "PRINTOUT_TYPE_RECEIPT": "DCC",
        "PRINTOUT_TYPE_TEMPLATE_DCC": "PLANET_PAX",
        "PRINTOUT_TYPE_DISPLAY": "Customer",
    }


@pytest.mark.parametrize("change_type", ["Receipt template DCC", "receipt TEMPLATE dcc"])
def test_parse_workbook_accepts_receipt_template_dcc(change_type):
    content = _workbook_bytes([
        [change_type, "I0001", "011276", "Shop", "Spain", "00123", "Planet", "T-123"]
    ])

    [target] = parse_workbook(content)

    assert target.check_column == "PRINTOUTTYPETEMPLATEDCC_CHECK_C"
    assert target.terminal_identifier == "T-123"


def test_parse_workbook_preserves_terminal_zeroes_from_excel_number_format():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    sheet.append(["dccEnable", "I0001", 11276, "Shop", "Spain", None, "Planet", 12345])
    sheet["H2"].number_format = "00000000"
    output = BytesIO()
    workbook.save(output)

    [target] = parse_workbook(output.getvalue())

    assert target.terminal_identifier == "00012345"
    assert target.metadata["LOCATION_NO"] == "11276"
    assert target.check_column == "HANDLER_DCCENABLE_CHECK_O"


def test_parse_workbook_accepts_instance_terminal_id_header():
    workbook = Workbook()
    sheet = workbook.active
    headers = ["InstanceTerminalID" if header == "TerminalID" else header for header in HEADERS]
    sheet.append(headers)
    sheet.append(
        ["handler_dccEnableNfc", "I0001", "011276", "Shop", "Spain", "00123", "Planet", "T-123"]
    )
    output = BytesIO()
    workbook.save(output)

    [target] = parse_workbook(output.getvalue())

    assert target.terminal_identifier == "T-123"
    assert target.check_column == "HANDLER_DCCENABLENFC_CHECK_O"


def test_parse_workbook_stops_at_first_fully_blank_row():
    content = _workbook_bytes(
        [
            ["handler_dccEnableNfc", "I0001", "11276", "Shop", "Spain", "00123", "Planet", "T-123"],
            [None] * len(HEADERS),
            ["handler_dccEnableNfc", "I0002", "11276", "Shop", "Spain", "00123", "Planet", None],
        ]
    )

    [target] = parse_workbook(content)

    assert target.terminal_identifier == "T-123"


def test_parse_workbook_rejects_unknown_check_and_missing_terminal_column():
    content = _workbook_bytes(
        [["unknown_flag", "I0001", "11276", "Shop", "Spain", "123", "Planet", "T1"]]
    )
    with pytest.raises(ValueError, match="unsupported Change Type"):
        parse_workbook(content)

    workbook = Workbook()
    workbook.active.append(["Anything", "location_no"])
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(ValueError, match="Missing required Excel column"):
        parse_workbook(output.getvalue())


def test_parse_workbook_rejects_unknown_metadata_header():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Check", "TerminalID", "Unknown field"])
    sheet.append(["dccEnable", "T-1", "value"])
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(ValueError, match="Unsupported Excel column"):
        parse_workbook(output.getvalue())


def test_missing_excel_dependency_does_not_crash_remediation(monkeypatch):
    monkeypatch.setattr(batch_upload, "load_workbook", None)
    with pytest.raises(ValueError, match="Excel upload requires openpyxl"):
        parse_workbook(b"not parsed because dependency is unavailable")


def test_uploaded_worklist_query_binds_terminal_ids_and_uses_only_fixable_checks():
    sql, params = wl.build_uploaded_targets_query(
        (
            ("T-001", "HANDLER_DCCENABLENFC_CHECK_O"),
            ("T-002", "CONFIGDOWNLOAD_VERSION_CHECK_C"),
        ),
        DEFAULT_OBJECTS,
    )

    assert "REMEDIATION_STATE = 'ACTIONABLE'" in sql
    assert "HANDLER_DCCENABLENFC_CHECK_O = 1" in sql
    assert "CONFIGDOWNLOAD_VERSION_CHECK_C = 1" in sql
    assert "T-001" not in sql and "T-002" not in sql
    assert params == ["T-001", "T-002"]
    with pytest.raises(ValueError, match="non-fixable"):
        wl.build_uploaded_targets_query((("T-001", "FIRMWARE_VERSION_CHECK"),), DEFAULT_OBJECTS)


class _FakeSnowflake:
    def __init__(self, row: dict):
        self.row = row
        self.query_sql = ""
        self.query_params = []

    def query(self, sql, params=None):
        self.query_sql = sql
        self.query_params = list(params or [])
        return pd.DataFrame([self.row])


def test_uploaded_rows_are_scoped_to_requested_broken_change_type_and_metadata():
    row = {
        "TERMINAL_IDENTIFIER": "T-001",
        "INSTANCE_IDENTIFIER": "I-001",
        "LOCATION_NO": "11276",
        "LOCATION_NAME": "Pilot ENG) Shop",
        "LOCATION_DESC": "Shop",
        "COUNTRY_NAME": "Spain",
        "BANK_MERCHANT_ID": "00123",
        "ACQUIRER_NAME": "Planet",
        "REMEDIATION_STATE": "ACTIONABLE",
        **{column: 0 for column in TRACKED_CHECK_COLUMNS},
        "HANDLER_DCCENABLENFC_CHECK_O": 1,
        "CONFIGDOWNLOAD_VERSION_CHECK_C": 1,
    }
    snow = _FakeSnowflake(row)
    content = _workbook_bytes(
        [["handler_dccEnableNfc", "I-001", "011276", "Shop", "Spain", "00123", "Planet", "T-001"]]
    )

    uploaded, warnings = _uploaded_batch_rows(snow, DEFAULT_OBJECTS, content)

    assert warnings == []
    assert snow.query_params == ["T-001"]
    assert "HANDLER_DCCENABLENFC_CHECK_O = 1" in snow.query_sql
    assert len(uploaded) == 1
    assert uploaded.iloc[0]["_BATCH_CHECK_COLUMNS"] == ("HANDLER_DCCENABLENFC_CHECK_O",)
    assert uploaded.iloc[0]["HANDLER_DCCENABLENFC_CHECK_O"] == 1
    assert uploaded.iloc[0]["CONFIGDOWNLOAD_VERSION_CHECK_C"] == 0


def test_uploaded_metadata_mismatch_skips_row():
    row = {
        "TERMINAL_IDENTIFIER": "T-001",
        "INSTANCE_IDENTIFIER": "I-001",
        "LOCATION_NO": "11276",
        "LOCATION_NAME": "Pilot ENG) Shop",
        "LOCATION_DESC": "Different description",
        "COUNTRY_NAME": "Spain",
        "BANK_MERCHANT_ID": "00123",
        "ACQUIRER_NAME": "Planet",
        "REMEDIATION_STATE": "ACTIONABLE",
        **{column: 0 for column in TRACKED_CHECK_COLUMNS},
        "HANDLER_DCCENABLENFC_CHECK_O": 1,
    }
    content = _workbook_bytes(
        [["handler_dccEnableNfc", "I-001", "011276", "Shop", "Spain", "00123", "Planet", "T-001"]]
    )

    uploaded, warnings = _uploaded_batch_rows(_FakeSnowflake(row), DEFAULT_OBJECTS, content)

    assert uploaded.empty
    assert any("LOCATION_DESC" in warning for warning in warnings)


def test_uploaded_provider_code_is_checked_against_snapshot():
    row = {
        "TERMINAL_IDENTIFIER": "T-001",
        "DCCPROVIDER_CODE": "PLANETPAYMENT",
        "REMEDIATION_STATE": "ACTIONABLE",
        **{column: 0 for column in TRACKED_CHECK_COLUMNS},
        "HANDLER_DCCENABLENFC_CHECK_O": 1,
    }
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Check", "TerminalID", "dccprovider_code"])
    sheet.append(["handler_dccEnableNfc", "T-001", "PLANETPAYMENT"])
    output = BytesIO()
    workbook.save(output)

    uploaded, warnings = _uploaded_batch_rows(
        _FakeSnowflake(row), DEFAULT_OBJECTS, output.getvalue()
    )

    assert warnings == []
    assert len(uploaded) == 1


def test_excel_batch_load_uses_uploaded_rows_only(monkeypatch):
    expected = pd.DataFrame([{"TERMINAL_IDENTIFIER": "T-UPLOAD"}])

    class Upload:
        def getvalue(self):
            return b"uploaded workbook"

    def upload_rows(conn, objs, content):
        assert content == b"uploaded workbook"
        return expected, []

    def filtered_rows(*_args, **_kwargs):
        raise AssertionError("Normal filters must not add rows when an Excel file is supplied.")

    monkeypatch.setattr("dcc_console.remediation.tab._uploaded_batch_rows", upload_rows)
    monkeypatch.setattr("dcc_console.remediation.tab._run_worklist", filtered_rows)

    actual, warnings = _load_batch_targets(
        object(), DEFAULT_OBJECTS, WorklistFilters(), Upload()
    )

    pd.testing.assert_frame_equal(actual, expected)
    assert warnings == []