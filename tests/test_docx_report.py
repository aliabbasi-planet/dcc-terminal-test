"""Tests for the signed Word (.docx) CAB report."""

from __future__ import annotations

import hashlib
import io

from docx import Document
from docx.oxml.ns import qn

from dcc_console.docx_report import generate_cab_docx
from dcc_console.execution import TestResult
from dcc_console.report import cab_report


def mk(**overrides) -> TestResult:
    defaults = dict(
        id="x", test_key="Bit 8 — DCC Handler Flags (instance)", bit=8,
        environment="DEV", login="svc", target_type="instance", target="I000000001",
        value="dccEnable", mode="SIMULATION", status="PASS", error=None,
        state_before="<x>1</x>", state_after="<x>1</x>", restore_point="<x>1</x>",
        change_persisted=False, transaction="ROLLED BACK (simulation)",
        sql="EXEC [cccai].[spApplyDCCEnablementConfiguration] @display_config = 8 ...",
        messages=["Trace: proposing dccEnable=true"], grids=[], grid_frames=[],
        duration_s=0.2, timestamp="2026-01-01T00:00:00",
        input_json='[{"instance_identifier":"I000000001"}]',
        proc_args=("@instance_json", "@Extra_Config_Name", "@Config_value", "@is_simulation"),
        proposed_change_observed=True, trace_status="CAPTURED",
        trace_rows=[{"line": 1, "message": "ok"}], return_code=0,
    )
    defaults.update(overrides)
    return TestResult(**defaults)


META = {"server": "LU3C01DVSQL01", "database": "3CDB", "environment": "DEV", "login": "svc-dcc"}


def _open(data: bytes) -> Document:
    return Document(io.BytesIO(data))


def test_generates_valid_docx_and_hash():
    data, content_hash = generate_cab_docx([mk()], dict(META))
    assert data[:2] == b"PK"  # docx is a zip
    assert len(content_hash) == 64
    doc = _open(data)
    assert len(doc.paragraphs) > 5


def test_hash_matches_canonical_markdown():
    # Fix generated_at so the digest is deterministic and independently reproducible.
    meta = dict(META, generated_at="2026-01-01T00:00:00+00:00")
    _, content_hash = generate_cab_docx([mk()], meta)
    # The hash is taken over the canonical markdown BEFORE the hash is embedded,
    # exactly like the PDF export — so reproduce it without content_hash in meta.
    expected_meta = dict(META, generated_at="2026-01-01T00:00:00+00:00")
    expected = hashlib.sha256(cab_report([mk()], expected_meta).encode("utf-8")).hexdigest()
    assert content_hash == expected


def test_renders_native_tables():
    data, _ = generate_cab_docx([mk()], dict(META))
    doc = _open(data)
    # Section B coverage + Section A evidence tables should be real Word tables.
    assert len(doc.tables) >= 2


def test_applies_read_only_protection():
    data, _ = generate_cab_docx([mk()], dict(META))
    doc = _open(data)
    protection = doc.settings.element.findall(qn("w:documentProtection"))
    assert protection, "documentProtection element missing"
    assert protection[0].get(qn("w:edit")) == "readOnly"
    assert protection[0].get(qn("w:enforcement")) == "1"


def test_manual_attachments_region_is_editable():
    data, _ = generate_cab_docx([mk()], dict(META))
    doc = _open(data)
    body_xml = doc.element.body.xml
    assert "Manual Attachments" in body_xml
    assert "permStart" in body_xml and "permEnd" in body_xml
    assert "everyone" in body_xml  # editable range group


def test_hash_embedded_in_core_properties():
    data, content_hash = generate_cab_docx([mk()], dict(META))
    doc = _open(data)
    assert content_hash in (doc.core_properties.keywords or "")
    assert content_hash in (doc.core_properties.comments or "")


def test_states_honest_integrity_limitation():
    data, _ = generate_cab_docx([mk()], dict(META))
    doc = _open(data)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "deterrent" in text
    assert "SHA-256" in text


# ---------------------------------------------------------------------------
# per-step evidence completeness (Bit 8 procedure-managed rollback)
# ---------------------------------------------------------------------------

_SCRIPT = (
    "UPDATE [cccintegrang].[handler] SET extra_config = "
    "'<extra_config config_name=\"dccEnableCompletion\" config_value=\"false\"/>' "
    "WHERE handler_id = 77"
)


def _all_text(doc: Document) -> str:
    """Every rendered string: paragraphs plus native table cells."""
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def mk_bit8_live(**overrides) -> TestResult:
    """A full Bit 8 live run: SP script captured, rolled back, flag verified."""
    return mk(
        mode="LIVE", status="PASS", value="dccEnableCompletion",
        transaction="COMMITTED then MANUAL ROLLED BACK (procedure script)",
        change_persisted=False,
        grids=[[{"rollback_script": _SCRIPT, "handler_name": "handler-A"}]],
        rollback_log=[{"ok": True, "rows": 1, "sql": _SCRIPT, "trigger": "manual",
                       "error": None, "at": "2026-09-21T20:48:00+00:00"}],
        sp_rollback_scripts=[_SCRIPT], sp_managed=True,
        sp_prior_states=[{"handler_name": "handler-A", "flag_value": "false"}],
        sp_after_states=[{"handler_name": "handler-A", "flag_value": "true"}],
        sp_restored_states=[{"handler_name": "handler-A", "flag_value": "false"}],
        sp_flag_verified=True,
        **overrides,
    )


def test_word_report_contains_every_required_evidence_item():
    """Guards the Word export against silently dropping any reported step."""
    data, _ = generate_cab_docx([mk_bit8_live()], dict(META))
    text = _all_text(_open(data))
    required = [
        "Procedure parameters",
        "Input JSON payload",
        "Database code exercised",
        "spApplyDCCEnablementConfiguration",       # 1 - the EXEC issued
        "Verification read",                        # 2 - verification SELECT
        "the procedure's OWN returned script",      # 3 - authoritative rollback
        "UPDATE [cccintegrang].[handler]",          # the SP script text itself
        "Rollback statements actually executed",    # 4 - what actually ran
        "Before / applied change / after evidence",
        "Procedure-managed bit",                    # honesty note
        "Handler-level verification",                # per-handler proof
        "After rollback",
        "Verified:",
        "handler-A",
        "Procedure return code",
        "Rollback verification",
        "Database message stream",
        "Procedure trace",
        "Procedure result set",
        "SHA-256",
        "deterrent",
        "Manual Attachments",
    ]
    for needle in required:
        assert needle in text, f"Word report is missing required evidence: {needle}"


def test_word_report_does_not_leak_markdown_markers():
    """A CAB document must not show raw '>', backticks or '_italics_' markers."""
    data, _ = generate_cab_docx([mk_bit8_live()], dict(META))
    paragraphs = [p.text for p in _open(data).paragraphs]
    assert not [t for t in paragraphs if t.strip().startswith(">")], "blockquote '>' leaked"
    assert not [t for t in paragraphs if "```" in t], "code fence leaked"
    assert not [t for t in paragraphs if "`" in t], "backtick leaked"
    italic = [
        t for t in paragraphs
        if len(t.strip()) > 2 and t.strip().startswith("_") and t.strip().endswith("_")
    ]
    assert not italic, f"italic markers leaked: {italic[:2]}"


def test_underscored_identifiers_survive_rendering():
    """Italic handling must never mangle names like extra_config."""
    text = _all_text(_open(generate_cab_docx([mk_bit8_live()], dict(META))[0]))
    assert "extra_config" in text
    assert "instance_identifier" in text
